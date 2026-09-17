from __future__ import annotations

import logging
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from value_eval.config import PACKAGE_ROOT, PipelineConfig
from value_eval.io_utils import atomic_write_json, load_json, sha256_file
from value_eval.value_to_scenario.excel_reader import read_and_validate_workbook
from value_eval.value_to_scenario.fingerprint import content_fingerprint
from value_eval.value_to_scenario.runner import PreparationOptions, ScenarioPreparationRunner


class ScenarioPreparationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.temp_root = Path(self.temporary.name)
        base = PipelineConfig.load(PACKAGE_ROOT / "configs" / "excel_to_benchmark.yaml")
        source = replace(
            base.scenario_source,
            input_xlsx=PACKAGE_ROOT / "inputs" / "excel" / "examples" / "value_first_2_rows.xlsx",
        )
        self.config = replace(
            base,
            output_root=self.temp_root / "outputs",
            log_root=self.temp_root / "logs",
            run_id="scenario_test",
            scenario_source=source,
        )

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _publish_fixture(self, *, method: str = "llm") -> dict:
        rows = read_and_validate_workbook(
            self.config.scenario_source.input_xlsx, self.config.scenario_source.sheet_name
        )
        template = load_json(
            PACKAGE_ROOT / "inputs/scenarios/examples/scenario_elements/example_fairness.json"
        )
        entries = []
        for row in (1, 2):
            for scenario in range(1, 8):
                scenario_id = f"row{row:03d}_s{scenario:02d}"
                path = self.config.prepared_scenario_dir / f"{scenario_id}.json"
                atomic_write_json(path, {**template, "scenario_id": scenario_id})
                entries.append({
                    "scenario_id": scenario_id,
                    "canonical_relative_path": path.name,
                    "sha256": sha256_file(path),
                    "generation_method": method,
                })
        manifest = {
            "schema_version": "value-to-scenario-manifest-v1",
            "fingerprint": content_fingerprint(self.config, rows),
            "valid": True,
            "published": True,
            "expected_scenarios": 14,
            "created_at": "2026-09-17T00:00:00Z",
            "entries": entries,
        }
        atomic_write_json(self.config.prepared_scenario_manifest, manifest)
        return manifest

    def test_resume_preserves_published_manifest_without_generation(self) -> None:
        manifest = self._publish_fixture()
        path = self.config.prepared_scenario_manifest
        original_bytes = path.read_bytes()
        original_mtime = path.stat().st_mtime_ns
        runner = ScenarioPreparationRunner(self.config, logging.getLogger(__name__))
        with patch(
            "value_eval.value_to_scenario.runner.prepare_taxonomy",
            side_effect=AssertionError("completed scenarios must not be regenerated"),
        ):
            for retry_fallbacks in (False, True):
                result = runner.run(PreparationOptions(retry_fallbacks=retry_fallbacks))
                self.assertEqual(result, manifest)
                self.assertEqual(path.read_bytes(), original_bytes)
                self.assertEqual(path.stat().st_mtime_ns, original_mtime)
        self.assertFalse((path.parent / "archive").exists())
        self.assertFalse((path.parent / "work").exists())

    def test_changed_or_incomplete_outputs_do_not_skip_preparation(self) -> None:
        for change in ("missing", "changed", "extra", "fingerprint", "invalid_json"):
            with self.subTest(change=change):
                self.config = replace(self.config, run_id=change)
                manifest = self._publish_fixture()
                path = self.config.prepared_scenario_dir / "row001_s01.json"
                if change == "missing":
                    path.unlink()
                elif change == "changed":
                    path.write_text("{}", encoding="utf-8")
                elif change == "extra":
                    (path.parent / "extra.json").write_text("{}", encoding="utf-8")
                elif change == "fingerprint":
                    manifest["fingerprint"] = "old-input"
                    atomic_write_json(self.config.prepared_scenario_manifest, manifest)
                else:
                    self.config.prepared_scenario_manifest.write_text("{", encoding="utf-8")
                self._assert_preparation_required(PreparationOptions())

    def test_force_no_publish_and_retry_fallbacks_are_respected(self) -> None:
        self._publish_fixture(method="heuristic_fallback")
        for options in (
            PreparationOptions(force=True),
            PreparationOptions(publish=False),
            PreparationOptions(retry_fallbacks=True),
        ):
            with self.subTest(options=options):
                self._assert_preparation_required(options)

    def _assert_preparation_required(self, options: PreparationOptions) -> None:
        runner = ScenarioPreparationRunner(self.config, logging.getLogger(__name__))
        with patch(
            "value_eval.value_to_scenario.runner.prepare_taxonomy",
            side_effect=RuntimeError("preparation reached"),
        ) as prepare:
            with self.assertRaisesRegex(RuntimeError, "preparation reached"):
                runner.run(options)
            prepare.assert_called_once()

    def test_excel_validation_reports_expected_scenarios(self) -> None:
        report = ScenarioPreparationRunner(self.config, logging.getLogger(__name__)).inspect_excel()
        self.assertTrue(report["valid"])
        self.assertEqual(report["workbook_rows"], 2)
        self.assertEqual(report["expected_scenarios"], 14)

    def test_excel_rejects_missing_header_and_empty_required_cell(self) -> None:
        headers = ["一级分类", "二级分类", "三级分类", "四级分类", "错误表头"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "价值观目录及评判标准"
        sheet.append(headers)
        sheet.append(["尊严", "维度", "子维度", "边界", "1. 场景"])
        invalid_header = self.temp_root / "invalid_header.xlsx"
        workbook.save(invalid_header)
        with self.assertRaisesRegex(ValueError, "缺少必要表头"):
            read_and_validate_workbook(invalid_header, sheet.title)

        sheet.cell(row=1, column=5, value="场景描述（评判标准）")
        sheet.cell(row=2, column=3, value="")
        empty_cell = self.temp_root / "empty_cell.xlsx"
        workbook.save(empty_cell)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "empty required fields"):
            read_and_validate_workbook(empty_cell, sheet.title)

if __name__ == "__main__":
    unittest.main()
