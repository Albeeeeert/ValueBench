from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from ..clients.openai_compat import OpenAICompatibleClient as LLMClient
from ..io_utils import extract_json_object
from .utils import (
    atomic_write_json,
    atomic_write_text,
    load_json_if_exists,
)

REQUIRED_HEADERS = ["一级分类", "二级分类", "三级分类", "四级分类", "场景描述（评判标准）"]
KNOWN_CATEGORY_SLUGS = {
    "Dignity": "dignity",
    "Stability": "stability",
    "Equality": "equality",
    "Freedom": "freedom",
    "Sustainability": "sustainability",
    "Inclusion": "inclusion",
    "Peace": "peace",
}
KNOWN_LEVEL1_TRANSLATIONS = {
    "尊严": "Dignity",
    "稳定": "Stability",
    "平等": "Equality",
    "自由": "Freedom",
    "可持续": "Sustainability",
    "包容": "Inclusion",
    "和平": "Peace",
}


@dataclass(frozen=True)
class SourceRow:
    row_id: int
    excel_row: int
    level1_zh: str
    level2_zh: str
    level3_zh: str
    level4_zh: str
    criteria_zh: str


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _slugify(text: str) -> str:
    text = text.strip()
    if text in KNOWN_CATEGORY_SLUGS:
        return KNOWN_CATEGORY_SLUGS[text]
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", text.lower()).strip("-")
    slug = re.sub(r"-{2,}", "-", slug)
    return slug or "unknown"


def _short_slug(text: str, max_len: int = 48) -> str:
    slug = _slugify(text)
    if len(slug) <= max_len:
        return slug
    digest = hashlib.sha1(text.encode("utf-8")).hexdigest()[:8]
    keep = max(12, max_len - len(digest) - 1)
    return f"{slug[:keep].rstrip('-')}-{digest}"


def _read_xlsx_rows(input_xlsx: Path, sheet_name: str) -> List[List[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(input_xlsx, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表: {sheet_name}; 可用工作表: {workbook.sheetnames}")
        sheet = workbook[sheet_name]
        return [[_as_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_source_rows(input_xlsx: Path, sheet_name: str, max_rows: Optional[int] = None) -> List[Dict[str, Any]]:
    if not input_xlsx.exists():
        raise FileNotFoundError(f"Excel 文件不存在: {input_xlsx}")

    rows = _read_xlsx_rows(input_xlsx, sheet_name)
    if not rows:
        raise ValueError(f"工作表为空: {sheet_name}")

    headers = [_as_text(v) for v in rows[0]]
    header_to_idx = {name: idx for idx, name in enumerate(headers) if name}
    missing = [h for h in REQUIRED_HEADERS if h not in header_to_idx]
    if missing:
        raise ValueError(f"Excel 缺少必要表头: {missing}; 实际表头: {headers}")

    out: List[Dict[str, Any]] = []
    for excel_row, raw in enumerate(rows[1:], start=2):
        values = {h: _as_text(raw[header_to_idx[h]] if header_to_idx[h] < len(raw) else "") for h in REQUIRED_HEADERS}
        if not any(values.values()):
            continue
        row = SourceRow(
            row_id=len(out) + 1,
            excel_row=excel_row,
            level1_zh=values["一级分类"],
            level2_zh=values["二级分类"],
            level3_zh=values["三级分类"],
            level4_zh=values["四级分类"],
            criteria_zh=values["场景描述（评判标准）"],
        )
        out.append(row.__dict__)
        if max_rows is not None and len(out) >= max_rows:
            break
    return out


def _chunk_rows(rows: List[Dict[str, Any]], chunk_by: str, chunk_size: int) -> List[List[Dict[str, Any]]]:
    size = max(1, int(chunk_size))
    if chunk_by == "level1":
        groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        order: List[str] = []
        for row in rows:
            key = str(row.get("level1_zh", ""))
            if key not in groups:
                order.append(key)
            groups[key].append(row)

        chunks: List[List[Dict[str, Any]]] = []
        for key in order:
            group = groups[key]
            for idx in range(0, len(group), size):
                chunks.append(group[idx : idx + size])
        return chunks

    chunks: List[List[Dict[str, Any]]] = []
    for idx in range(0, len(rows), size):
        chunks.append(rows[idx : idx + size])
    return chunks


def _chat_json(
    *,
    client: LLMClient,
    system_prompt: str,
    user_prompt: str,
    save_prefix: Optional[Path] = None,
) -> Dict[str, Any]:
    if save_prefix is not None:
        atomic_write_text(save_prefix.with_suffix(".system.txt"), system_prompt + "\n")
        atomic_write_text(save_prefix.with_suffix(".prompt.txt"), user_prompt + "\n")

    raw = client.chat_completions(
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )

    if save_prefix is not None:
        atomic_write_text(save_prefix.with_suffix(".raw.txt"), raw if raw.endswith("\n") else raw + "\n")

    try:
        parsed = extract_json_object(raw)
    except Exception as exc:
        if save_prefix is not None:
            atomic_write_text(
                save_prefix.with_suffix(".error.txt"),
                f"parse_error: {repr(exc)}\nraw_chars={len(raw)}\nraw_tail:\n{raw[-2000:]}\n",
            )
        raise
    if not isinstance(parsed, dict):
        raise ValueError("LLM 输出不是 JSON object")
    if save_prefix is not None:
        atomic_write_json(save_prefix.with_suffix(".parsed.json"), parsed)
    return parsed


def _translation_system_prompt() -> str:
    return (
        "You are a careful Chinese-to-English translator for AI safety benchmark taxonomies. "
        "Return strict JSON only. Translate faithfully into clear English. Do not summarize, reinterpret, "
        "or add new policy content. Preserve row_id values. Output English in translated fields; keep any "
        "explicit source-trace fields unchanged if requested."
    )


def _translation_user_prompt(rows: List[Dict[str, Any]]) -> str:
    return (
        "Translate the following value taxonomy rows into English.\n"
        "Return exactly this JSON schema: {\"rows\": [{\"row_id\": number, \"level1_en\": string, "
        "\"level2_en\": string, \"level3_en\": string, \"level4_en\": string, \"criteria_en\": string, "
        "\"translation_notes\": string}]}.\n"
        "Rules: criteria_en should be a faithful readable translation of the Chinese judging criteria. "
        "Do not expand unsafe criteria into additional operational instructions.\n\n"
        f"Rows:\n{json.dumps(rows, ensure_ascii=False, indent=2)}"
    )


def _extraction_system_prompt() -> str:
    return (
        "You are an AI safety benchmark taxonomy analyst. Return strict JSON only. Convert translated value "
        "taxonomy rows into benchmark-ready English taxonomy dimensions and high-level judging criteria. "
        "Do not create actionable harmful instructions. For unsafe domains, phrase criteria as evaluation "
        "boundaries and content-risk indicators. Preserve row_id and category paths."
    )


def _extraction_user_prompt(rows: List[Dict[str, Any]]) -> str:
    return (
        "Extract benchmark-ready value paradigms from these translated rows.\n"
        "Return exactly this JSON schema: {\"records\": [{\"row_id\": number, "
        "\"category_path_en\": [string, string, string, string], \"value_paradigm\": string, "
        "\"risk_domain\": string, \"benchmark_dimension\": string, \"judging_criteria\": [string], "
        "\"safe_content_boundary\": string, \"benchmark_generation_notes\": string}]}.\n"
        "Rules: judging_criteria must be concise English evaluation criteria, not procedural how-to steps. "
        "safe_content_boundary should describe allowed benign/educational/prevention framing.\n\n"
        f"Translated rows:\n{json.dumps(rows, ensure_ascii=False, indent=2)}"
    )


def _merge_translation(source_rows: List[Dict[str, Any]], translated_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_id = {int(r.get("row_id")): r for r in translated_rows if isinstance(r, dict) and r.get("row_id") is not None}
    merged: List[Dict[str, Any]] = []
    for src in source_rows:
        row_id = int(src["row_id"])
        tr = by_id.get(row_id, {})
        merged.append({
            **src,
            "level1_en": str(tr.get("level1_en") or KNOWN_LEVEL1_TRANSLATIONS.get(str(src.get("level1_zh")), "")).strip(),
            "level2_en": str(tr.get("level2_en", "")).strip(),
            "level3_en": str(tr.get("level3_en", "")).strip(),
            "level4_en": str(tr.get("level4_en", "")).strip(),
            "criteria_en": str(tr.get("criteria_en", "")).strip(),
            "translation_notes": str(tr.get("translation_notes", "")).strip(),
        })
    return merged


def _escape_md(text: Any, max_len: Optional[int] = None) -> str:
    s = str(text or "").replace("\n", " ").replace("|", "\\|").strip()
    if max_len and len(s) > max_len:
        return s[: max_len - 3] + "..."
    return s


def _split_numbered_sections(text: Any) -> List[Dict[str, Any]]:
    raw = str(text or "").strip()
    if not raw:
        return []

    pattern = re.compile(r"(?:^|\n)\s*(\d+)[\.．、]\s*")
    matches = list(pattern.finditer(raw))
    if not matches:
        return [{"scene_no": 1, "text": raw}]

    sections: List[Dict[str, Any]] = []
    prefix = raw[: matches[0].start()].strip()
    for idx, match in enumerate(matches):
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(raw)
        body = raw[start:end].strip()
        if prefix and idx == 0:
            body = f"{prefix}\n{body}".strip()
        if not body:
            continue
        try:
            scene_no = int(match.group(1))
        except Exception:
            scene_no = len(sections) + 1
        sections.append({"scene_no": scene_no, "text": body})
    return sections or [{"scene_no": 1, "text": raw}]


def _title_from_scene(text: Any) -> str:
    s = str(text or "").strip().replace("\n", " ")
    if not s:
        return "Scenario"
    for sep in [":", "："]:
        if sep in s:
            head = s.split(sep, 1)[0].strip()
            if head:
                return head[:160]
    return s[:160]


def _get_record_by_id(records: Optional[Sequence[Dict[str, Any]]]) -> Dict[int, Dict[str, Any]]:
    if not records:
        return {}
    out: Dict[int, Dict[str, Any]] = {}
    for rec in records:
        try:
            out[int(rec.get("row_id"))] = rec
        except Exception:
            continue
    return out


def write_translation_markdown(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    lines = [
        "# Value Taxonomy Translation",
        "",
        "| Row | Level 1 | Level 2 | Level 3 | Level 4 | Criteria |",
        "|---:|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| {row_id} | {l1} | {l2} | {l3} | {l4} | {criteria} |".format(
                row_id=row.get("row_id"),
                l1=_escape_md(row.get("level1_en")),
                l2=_escape_md(row.get("level2_en")),
                l3=_escape_md(row.get("level3_en")),
                l4=_escape_md(row.get("level4_en")),
                criteria=_escape_md(row.get("criteria_en"), 1200),
            )
        )
    atomic_write_text(path, "\n".join(lines) + "\n")


def translate_rows(
    *,
    rows: List[Dict[str, Any]],
    client: LLMClient,
    output_root: Path,
    chunk_by: str,
    chunk_size: int,
    debug_save_llm_io: bool,
    resume: bool,
) -> List[Dict[str, Any]]:
    out_path = output_root / "translation" / "translated_rows.en.json"
    if resume and out_path.exists():
        existing = load_json_if_exists(out_path, [])
        if isinstance(existing, list) and len(existing) >= len(rows):
            client.logger.info("复用已有翻译: %s", out_path)
            return existing[: len(rows)]

    translated: List[Dict[str, Any]] = []
    chunks = _chunk_rows(rows, chunk_by, chunk_size)
    for idx, chunk in enumerate(chunks, start=1):
        label = str(chunk[0].get("level1_zh") or idx)
        save_prefix = output_root / "logs" / "llm_raw" / f"translate_{idx:03d}_{_slugify(label)}" if debug_save_llm_io else None
        parsed = _chat_json(
            client=client,
            system_prompt=_translation_system_prompt(),
            user_prompt=_translation_user_prompt(chunk),
            save_prefix=save_prefix,
        )
        chunk_rows = parsed.get("rows", [])
        if not isinstance(chunk_rows, list):
            raise ValueError("翻译 LLM 输出缺少 rows 列表")
        translated.extend(chunk_rows)
        client.logger.info("翻译完成 chunk %d/%d | rows=%d", idx, len(chunks), len(chunk_rows))

    merged = _merge_translation(rows, translated)
    atomic_write_json(out_path, merged)
    write_translation_markdown(output_root / "translation" / "translated_table.en.md", merged)
    return merged


def extract_records(
    *,
    translated_rows: List[Dict[str, Any]],
    client: LLMClient,
    output_root: Path,
    chunk_by: str,
    chunk_size: int,
    debug_save_llm_io: bool,
    resume: bool,
) -> List[Dict[str, Any]]:
    out_path = output_root / "extracted" / "review_criteria.en.json"
    if resume and out_path.exists():
        existing = load_json_if_exists(out_path, [])
        if isinstance(existing, list) and len(existing) >= len(translated_rows):
            client.logger.info("复用已有抽取结果: %s", out_path)
            return existing[: len(translated_rows)]

    records: List[Dict[str, Any]] = []
    chunks = _chunk_rows(translated_rows, chunk_by, chunk_size)
    for idx, chunk in enumerate(chunks, start=1):
        label = str(chunk[0].get("level1_en") or chunk[0].get("level1_zh") or idx)
        save_prefix = output_root / "logs" / "llm_raw" / f"extract_{idx:03d}_{_slugify(label)}" if debug_save_llm_io else None
        parsed = _chat_json(
            client=client,
            system_prompt=_extraction_system_prompt(),
            user_prompt=_extraction_user_prompt(chunk),
            save_prefix=save_prefix,
        )
        chunk_records = parsed.get("records", [])
        if not isinstance(chunk_records, list):
            raise ValueError("抽取 LLM 输出缺少 records 列表")
        records.extend(chunk_records)
        client.logger.info("抽取完成 chunk %d/%d | records=%d", idx, len(chunks), len(chunk_records))

    merged = _merge_extracted(translated_rows, records)
    atomic_write_json(out_path, merged)
    return merged


def _merge_extracted(translated_rows: List[Dict[str, Any]], records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_id = {int(r.get("row_id")): r for r in records if isinstance(r, dict) and r.get("row_id") is not None}
    out: List[Dict[str, Any]] = []
    for row in translated_rows:
        row_id = int(row["row_id"])
        rec = by_id.get(row_id, {})
        category_path = rec.get("category_path_en")
        if not isinstance(category_path, list) or len(category_path) != 4:
            category_path = [row.get("level1_en", ""), row.get("level2_en", ""), row.get("level3_en", ""), row.get("level4_en", "")]
        criteria = rec.get("judging_criteria", [])
        if not isinstance(criteria, list):
            criteria = [str(criteria)] if str(criteria).strip() else []
        out.append({
            "row_id": row_id,
            "source": {
                "excel_row": row.get("excel_row"),
                "category_path_zh": [row.get("level1_zh"), row.get("level2_zh"), row.get("level3_zh"), row.get("level4_zh")],
                "criteria_zh": row.get("criteria_zh"),
            },
            "category_path_en": [str(x).strip() for x in category_path],
            "criteria_en": row.get("criteria_en", ""),
            "value_paradigm": str(rec.get("value_paradigm", "")).strip(),
            "risk_domain": str(rec.get("risk_domain", "")).strip(),
            "benchmark_dimension": str(rec.get("benchmark_dimension", "")).strip(),
            "judging_criteria": [str(x).strip() for x in criteria if str(x).strip()],
            "safe_content_boundary": str(rec.get("safe_content_boundary", "")).strip(),
            "benchmark_generation_notes": str(rec.get("benchmark_generation_notes", "")).strip(),
        })
    return out


def build_taxonomy(records: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    categories: Dict[str, Dict[str, Any]] = {}
    for rec in records:
        path = [str(x or "").strip() for x in rec.get("category_path_en", [])]
        if not path or not path[0]:
            path = ["Unknown", "", "", ""]
        level1 = path[0]
        cat = categories.setdefault(level1, {"theme": level1, "slug": _slugify(level1), "subdimensions": []})
        cat["subdimensions"].append({
            "row_id": rec.get("row_id"),
            "name": path[3] if len(path) > 3 and path[3] else (path[-1] if path else ""),
            "path": path,
            "description": rec.get("benchmark_dimension") or rec.get("risk_domain") or rec.get("criteria_en", ""),
            "value_paradigm": rec.get("value_paradigm", ""),
            "risk_domain": rec.get("risk_domain", ""),
            "judging_criteria": rec.get("judging_criteria", []),
            "safe_content_boundary": rec.get("safe_content_boundary", ""),
            "source": rec.get("source", {}),
        })
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "language": "en",
        "categories": list(categories.values()),
    }


def write_benchmark_paradigms(path: Path, records: Sequence[Dict[str, Any]]) -> None:
    by_paradigm: Dict[str, Dict[str, Any]] = {}
    for rec in records:
        key = str(rec.get("value_paradigm") or "Unspecified").strip()
        node = by_paradigm.setdefault(key, {"value_paradigm": key, "dimensions": []})
        node["dimensions"].append({
            "row_id": rec.get("row_id"),
            "category_path_en": rec.get("category_path_en", []),
            "risk_domain": rec.get("risk_domain", ""),
            "benchmark_dimension": rec.get("benchmark_dimension", ""),
        })
    atomic_write_json(path, list(by_paradigm.values()))


def write_category_outputs(output_root: Path, taxonomy: Dict[str, Any]) -> None:
    for idx, category in enumerate(taxonomy.get("categories", []), start=1):
        theme = str(category.get("theme", f"category-{idx}"))
        slug = str(category.get("slug") or _slugify(theme))
        cat_dir = output_root / "by_category" / slug
        payload = {
            "id": idx,
            "theme": theme,
            "slug": slug,
            "language": "en",
            "subdimensions": category.get("subdimensions", []),
        }
        atomic_write_json(cat_dir / "taxonomy.json", payload)
        write_category_markdown(cat_dir / "criteria.md", payload)


def write_category_markdown(path: Path, category: Dict[str, Any]) -> None:
    lines = [f"# {category.get('theme', '')}", "", "| Row | Path | Paradigm | Criteria |", "|---:|---|---|---|"]
    for item in category.get("subdimensions", []):
        criteria = "; ".join(str(x) for x in item.get("judging_criteria", []))
        lines.append(
            "| {row_id} | {path} | {paradigm} | {criteria} |".format(
                row_id=item.get("row_id", ""),
                path=_escape_md(" > ".join(str(x) for x in item.get("path", []))),
                paradigm=_escape_md(item.get("value_paradigm", "")),
                criteria=_escape_md(criteria, 1000),
            )
        )
    atomic_write_text(path, "\n".join(lines) + "\n")


def _dimension_dir_for_row(output_root: Path, row: Dict[str, Any]) -> Path:
    parts = [
        row.get("level1_en") or KNOWN_LEVEL1_TRANSLATIONS.get(str(row.get("level1_zh", "")), "") or row.get("level1_zh", "unknown"),
        row.get("level2_en") or row.get("level2_zh", "unknown"),
        row.get("level3_en") or row.get("level3_zh", "unknown"),
        row.get("level4_en") or row.get("level4_zh", "unknown"),
    ]
    return output_root / "by_dimension" / Path(*[_short_slug(str(p)) for p in parts])


def _category_path_en(row: Dict[str, Any]) -> List[str]:
    return [
        str(row.get("level1_en") or KNOWN_LEVEL1_TRANSLATIONS.get(str(row.get("level1_zh", "")), "")).strip(),
        str(row.get("level2_en", "")).strip(),
        str(row.get("level3_en", "")).strip(),
        str(row.get("level4_en", "")).strip(),
    ]


def _category_path_zh(row: Dict[str, Any]) -> List[str]:
    return [
        str(row.get("level1_zh", "")).strip(),
        str(row.get("level2_zh", "")).strip(),
        str(row.get("level3_zh", "")).strip(),
        str(row.get("level4_zh", "")).strip(),
    ]


def build_dimension_payload(row: Dict[str, Any], record: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    record = record or {}
    source_sections = _split_numbered_sections(row.get("criteria_zh", ""))
    translation_sections = _split_numbered_sections(row.get("criteria_en", ""))
    criteria = record.get("judging_criteria", [])
    if not isinstance(criteria, list):
        criteria = [str(criteria)] if str(criteria).strip() else []

    scenarios: List[Dict[str, Any]] = []
    for idx, section in enumerate(source_sections, start=1):
        translated = translation_sections[idx - 1]["text"] if idx - 1 < len(translation_sections) else ""
        scene_no = int(section.get("scene_no") or idx)
        scenarios.append({
            "scenario_id": f"row{int(row.get('row_id', 0)):03d}_s{scene_no:02d}",
            "scene_no": scene_no,
            "scenario_title_en": _title_from_scene(translated) if translated else "",
            "scenario_title_zh": _title_from_scene(section.get("text", "")),
            "source_text_zh": section.get("text", ""),
            "translation_en": translated,
            "judging_criteria_en": criteria,
            "safe_content_boundary": record.get("safe_content_boundary", ""),
            "benchmark_generation_notes": record.get("benchmark_generation_notes", ""),
        })

    return {
        "row_id": row.get("row_id"),
        "excel_row": row.get("excel_row"),
        "language": "en",
        "category_path_zh": _category_path_zh(row),
        "category_path_en": _category_path_en(row),
        "value_paradigm": record.get("value_paradigm", ""),
        "risk_domain": record.get("risk_domain", ""),
        "benchmark_dimension": record.get("benchmark_dimension", ""),
        "criteria_zh_full": row.get("criteria_zh", ""),
        "criteria_en_full": row.get("criteria_en", ""),
        "scenarios": scenarios,
    }


def write_hierarchical_dimension_outputs(
    output_root: Path,
    rows: Sequence[Dict[str, Any]],
    records: Optional[Sequence[Dict[str, Any]]] = None,
) -> None:
    record_by_id = _get_record_by_id(records)
    index: List[Dict[str, Any]] = []
    scenario_index: List[Dict[str, Any]] = []
    used_paths: Dict[str, int] = defaultdict(int)

    for row in rows:
        try:
            row_id = int(row.get("row_id"))
        except Exception:
            continue
        dim_dir = _dimension_dir_for_row(output_root, row)
        path_key = str(dim_dir)
        used_paths[path_key] += 1
        if used_paths[path_key] > 1:
            dim_dir = dim_dir.with_name(f"{dim_dir.name}-row{row_id:03d}")

        payload = build_dimension_payload(row, record_by_id.get(row_id))
        rel_dir = dim_dir.relative_to(output_root).as_posix()
        atomic_write_json(dim_dir / "scenarios.json", payload)
        index.append({
            "row_id": row_id,
            "dimension_dir": rel_dir,
            "category_path_zh": payload["category_path_zh"],
            "category_path_en": payload["category_path_en"],
            "scenario_count": len(payload["scenarios"]),
        })
        for scenario in payload["scenarios"]:
            scenario_index.append({
                "scenario_id": scenario.get("scenario_id"),
                "row_id": row_id,
                "dimension_dir": rel_dir,
                "category_path_en": payload["category_path_en"],
                "scenario_title_en": scenario.get("scenario_title_en", ""),
            })

    atomic_write_json(output_root / "dimension_index.en.json", index)
    atomic_write_json(output_root / "scenario_index.en.json", scenario_index)


def write_extracted_outputs(output_root: Path, records: List[Dict[str, Any]], translated_rows: Optional[List[Dict[str, Any]]] = None) -> None:
    taxonomy = build_taxonomy(records)
    atomic_write_json(output_root / "extracted" / "taxonomy.en.json", taxonomy)
    write_benchmark_paradigms(output_root / "extracted" / "benchmark_paradigms.en.json", records)
    atomic_write_json(output_root / "extracted" / "review_criteria.en.json", records)
    write_category_outputs(output_root, taxonomy)
    if translated_rows is not None:
        write_hierarchical_dimension_outputs(output_root, translated_rows, records)


def load_translated_rows(output_root: Path) -> List[Dict[str, Any]]:
    path = output_root / "translation" / "translated_rows.en.json"
    data = load_json_if_exists(path, [])
    if not isinstance(data, list) or not data:
        raise ValueError(f"缺少翻译结果，请先运行翻译阶段: {path}")
    return data
